"""Turn uploaded files into labelled chunks of text (one per page or slide)."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import PurePath

SUPPORTED = (".pdf", ".pptx", ".txt", ".md")
# Plain-text files have no pages, so split them into page-sized pieces.
TEXT_PAGE_CHARS = 3000
# A PDF page with less text than this is mostly a picture (a scan, a diagram, a slide
# exported as an image), so Claude reads the page image during the build instead.
IMAGE_PAGE_CHARS = 40


@dataclass
class Chunk:
    label: str
    text: str
    page: int = 0


@dataclass
class Extracted:
    chunks: list[Chunk]
    image_pages: list[int] = field(default_factory=list)  # 1-based PDF pages for Claude to read


class UnsupportedFile(ValueError):
    pass


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract(filename: str, data: bytes) -> Extracted:
    ext = PurePath(filename).suffix.lower()
    if ext == ".pdf":
        return _pdf(filename, data)
    if ext == ".pptx":
        return Extracted(_pptx(filename, data))
    if ext in (".txt", ".md"):
        return Extracted(_text(filename, data.decode("utf-8", errors="replace")))
    raise UnsupportedFile(f"{filename}: only PDF, PowerPoint (.pptx), .txt and .md files are supported")


def pdf_label(filename: str, page: int) -> str:
    return f"{filename}, p.{page}"


def _pdf(filename: str, data: bytes) -> Extracted:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    out = Extracted([])
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = _clean(page.extract_text() or "")
        except Exception:
            text = ""
        if len(text) >= IMAGE_PAGE_CHARS:
            out.chunks.append(Chunk(pdf_label(filename, i), text, i))
        else:
            out.image_pages.append(i)
    return out


def pdf_pages(data: bytes, pages: list[int]) -> bytes:
    """A new PDF holding just these 1-based pages, to send to Claude."""
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(io.BytesIO(data))
    writer = PdfWriter()
    for p in pages:
        writer.add_page(reader.pages[p - 1])
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _pptx(filename: str, data: bytes) -> list[Chunk]:
    from pptx import Presentation

    deck = Presentation(io.BytesIO(data))
    chunks = []
    for i, slide in enumerate(deck.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            elif getattr(shape, "has_table", False) and shape.has_table:
                for r in shape.table.rows:
                    parts.append(" | ".join(c.text for c in r.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Speaker notes: {notes}")
        text = _clean("\n".join(parts))
        if text:
            chunks.append(Chunk(f"{filename}, slide {i}", text, i))
    return chunks


def _text(filename: str, text: str) -> list[Chunk]:
    text = _clean(text)
    chunks, part, n = [], [], 0
    for para in text.split("\n\n"):
        if part and n + len(para) > TEXT_PAGE_CHARS:
            chunks.append("\n\n".join(part))
            part, n = [], 0
        part.append(para)
        n += len(para)
    if part:
        chunks.append("\n\n".join(part))
    if len(chunks) == 1:
        return [Chunk(filename, chunks[0], 1)] if chunks[0] else []
    return [Chunk(f"{filename}, part {i}", c, i) for i, c in enumerate(chunks, start=1)]
