import io

import pytest

from lectern.extract import UnsupportedFile, extract


def test_text_splits_into_parts():
    text = "\n\n".join(f"Paragraph {i} " + "x" * 900 for i in range(10))
    chunks = extract("notes.md", text.encode())
    assert len(chunks) > 1 and chunks[0].label == "notes.md, part 1"
    assert "Paragraph 0" in chunks[0].text


def test_short_text_is_one_chunk():
    chunks = extract("notes.txt", b"Just one idea.")
    assert [(c.label, c.text) for c in chunks] == [("notes.txt", "Just one idea.")]


def test_pptx_slides_and_notes():
    from pptx import Presentation

    deck = Presentation()
    for title in ("Limits", ""):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
    deck.slides[0].notes_slide.notes_text_frame.text = "Mention epsilon-delta"
    buf = io.BytesIO()
    deck.save(buf)
    chunks = extract("week1.pptx", buf.getvalue())
    assert len(chunks) == 1  # the empty slide is skipped
    assert chunks[0].label == "week1.pptx, slide 1"
    assert "Limits" in chunks[0].text and "epsilon-delta" in chunks[0].text


def test_pdf_pages():
    stream = b"BT /F1 12 Tf 72 720 Td (Chain rule) Tj ET"
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % off for off in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    chunks = extract("book.pdf", out)
    assert [(c.label, c.text) for c in chunks] == [("book.pdf, p.1", "Chain rule")]


def test_unsupported():
    with pytest.raises(UnsupportedFile):
        extract("slides.key", b"")
